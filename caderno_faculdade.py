# -*- coding: utf-8 -*-
"""
Caderno Faculdade
-----------------
Caderno digital que guarda cada matéria como um arquivo .docx dentro de
    <pasta>/<Período>/<Disciplina>/<Título>.docx
Se a pasta estiver dentro do Google Drive para computador, o Drive sincroniza.

Novidades desta versão
  * Salvamento automático (com gravação segura, sem arquivos corrompidos no Drive)
  * Detecção de alterações feitas pelo Drive em outro computador
  * Painel inicial com estatísticas, atividade semanal e matérias recentes
  * Nova interface (barra lateral escura, cartões arredondados, tema claro/escuro)
  * Atalhos: Ctrl+S salvar agora, Ctrl+N nova matéria, Ctrl+K buscar
"""
import json
import os
import re
import shutil
import subprocess
import sys
import threading
import time
from datetime import datetime, timedelta
from pathlib import Path

import tkinter as tk
import tkinter.font as tkfont
from tkinter import ttk, filedialog, messagebox

from docx import Document
from docx.shared import Pt, RGBColor

from ui_render import rounded_rect_photo, switch_photo

APP_NAME = "Caderno Faculdade"
CONFIG_DIR = Path.home() / ".caderno_faculdade"
CONFIG_FILE = CONFIG_DIR / "config.json"
DEFAULT_ROOT = Path.home() / "Caderno Faculdade"

ALL_PERIODS = "Todos os períodos"
ALL_COURSES = "Todas as disciplinas"

AUTOSAVE_BODY_MS = 1500      # pausa na digitação do texto antes de salvar
AUTOSAVE_META_MS = 3500      # pausa ao editar título/período/disciplina (evita renomear a cada letra)
AUTOSAVE_MAX_WAIT_MS = 12000 # nunca passa mais que isso sem salvar enquanto digita
RESCAN_EVERY_TICKS = 4       # a cada 4 ticks (4s cada) relê a pasta
TICK_MS = 4000

# ----------------------------------------------------------------------------
# Paleta (a barra lateral é escura nos dois temas, como na referência)
# ----------------------------------------------------------------------------
_SIDE = dict(
    side="#0b1510", side2="#15241b", side_hover="#1d3227", side_fg="#e8f1eb",
    side_muted="#7d9185", side_line="#24382c",
)
THEMES = {
    "light": dict(
        _SIDE,
        bg="#edf2ed", card="#ffffff", fg="#0f1a13", muted="#6a7a70", line="#e2e9e3",
        field="#f2f6f2", accent="#2eb54f", accent_hover="#27a047", accent_fg="#ffffff",
        accent_text="#1f8f3c", soft="#dcebdc", soft_fg="#10301b",
        dark="#0e1b13", dark_fg="#ffffff", dark_muted="#8fa699", dark_hover="#1a2e22",
        bar="#cfe0d2", danger="#d64545", warn="#e8a317", select="#cfe8d4",
    ),
    "dark": dict(
        _SIDE,
        bg="#0e1712", card="#16231b", fg="#e8f0ea", muted="#8a9d92", line="#25362c",
        field="#101b14", accent="#3fd06b", accent_hover="#35bb5d", accent_fg="#06130b",
        accent_text="#4be07a", soft="#1c3326", soft_fg="#e4f3e8",
        dark="#09110c", dark_fg="#ffffff", dark_muted="#7f9488", dark_hover="#13221a",
        bar="#2a3d31", danger="#ef6a6a", warn="#f0b33a", select="#254632",
    ),
}

WEEKDAYS = ["seg", "ter", "qua", "qui", "sex", "sáb", "dom"]


# ----------------------------------------------------------------------------
# Funções utilitárias (sem interface — fáceis de testar)
# ----------------------------------------------------------------------------
def clean(value, default):
    """Nome seguro para pasta/arquivo (Windows incluído)."""
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "-", (value or "").strip())
    value = re.sub(r"\s+", " ", value).strip().rstrip(". ")[:100].strip()
    if value.upper() in {"CON", "PRN", "AUX", "NUL"} or re.fullmatch(r"(COM|LPT)\d", value.upper() or ""):
        value += "_"
    return value or default


def first_line(text, limit=60):
    for line in text.splitlines():
        line = line.strip()
        if line:
            return line[:limit]
    return ""


def unique_path(path):
    if not path.exists():
        return path
    i = 2
    while True:
        cand = path.with_name(f"{path.stem} ({i}){path.suffix}")
        if not cand.exists():
            return cand
        i += 1


def same_path(a, b):
    return os.path.normcase(os.path.abspath(str(a))) == os.path.normcase(os.path.abspath(str(b)))


def trunc(text, n):
    text = text or ""
    return text if len(text) <= n else text[: n - 1].rstrip() + "…"


def rel_time(ts):
    d = datetime.fromtimestamp(ts)
    n = datetime.now()
    if d.date() == n.date():
        return f"Hoje, {d:%H:%M}"
    if d.date() == (n - timedelta(days=1)).date():
        return f"Ontem, {d:%H:%M}"
    if 0 <= (n - d).days < 7:
        return f"{WEEKDAYS[d.weekday()].capitalize()}, {d:%H:%M}"
    return d.strftime("%d/%m/%Y")


def is_drive_path(path):
    s = str(path).lower()
    return any(t in s for t in ("google drive", "googledrive", "meu drive", "my drive", "drive.google"))


def count_words(text):
    return len(re.findall(r"\w+", text))


XML_BAD = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")


def write_docx(path, title, period, course, content):
    """Grava o DOCX de forma segura: escreve em .tmp e só então substitui o arquivo."""
    doc = Document()
    try:
        normal = doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(6)
        h1 = doc.styles["Heading 1"]
        h1.font.name = "Calibri"
        h1.font.color.rgb = RGBColor(0x1F, 0x7A, 0x3A)
    except Exception:
        pass
    doc.add_heading(XML_BAD.sub("", title), level=1)
    meta = doc.add_paragraph()
    run = meta.add_run(f"Período: {period}  |  Disciplina: {course}")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x6A, 0x7A, 0x70)
    doc.add_paragraph("")
    for block in XML_BAD.sub("", content).split("\n\n"):
        doc.add_paragraph(block)
    tmp = path.with_name(path.name + ".tmp")
    doc.save(str(tmp))
    os.replace(str(tmp), str(path))


def extract_content(doc):
    """Texto da matéria, sem o cabeçalho (título + linha de período) criado pelo app."""
    paras = list(doc.paragraphs)
    i = 0
    if len(paras) >= 2:
        style = str(getattr(paras[0].style, "name", "")).lower()
        if style.startswith(("heading", "título", "title")) and paras[1].text.startswith("Período:"):
            i = 2
            if i < len(paras) and not paras[i].text.strip():
                i += 1
    return "\n\n".join(p.text for p in paras[i:])


def read_note(path):
    return extract_content(Document(str(path)))


def read_config():
    try:
        return json.loads(CONFIG_FILE.read_text(encoding="utf-8"))
    except Exception:
        return {}


def write_config(data):
    try:
        CONFIG_DIR.mkdir(parents=True, exist_ok=True)
        CONFIG_FILE.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    except OSError:
        pass


def rounded_points(x1, y1, x2, y2, r):
    r = max(0, min(r, (x2 - x1) / 2, (y2 - y1) / 2))
    return [x1 + r, y1, x1 + r, y1, x2 - r, y1, x2 - r, y1, x2, y1, x2, y1 + r, x2, y1 + r,
            x2, y2 - r, x2, y2 - r, x2, y2, x2 - r, y2, x2 - r, y2, x1 + r, y2, x1 + r, y2,
            x1, y2, x1, y2 - r, x1, y2 - r, x1, y1 + r, x1, y1 + r, x1, y1]


# ----------------------------------------------------------------------------
# Widgets personalizados
# (atributos com "_" na frente para não colidir com os do Tkinter)
# ----------------------------------------------------------------------------
class Surface(tk.Canvas):
    """Cartão de cantos arredondados. Coloque o conteúdo em `.inner`."""

    def __init__(self, parent, app, fill="card", on="bg", radius=18, pad=16, width=10, height=10):
        super().__init__(parent, width=width, height=height, highlightthickness=0, bd=0)
        self._app, self._fill, self._on, self._radius, self._pad = app, fill, on, radius, pad
        self._cw = self._ch = 0
        self.inner = tk.Frame(self, bd=0, highlightthickness=0)
        self._win = self.create_window(pad, pad, window=self.inner, anchor="nw")
        self.bind("<Configure>", self._on_configure)
        app.register(self)
        self.restyle()

    def restyle(self):
        self.configure(bg=self._app.c(self._on))
        self.inner.configure(bg=self._app.c(self._fill))
        self._redraw()

    def _on_configure(self, e):
        self._cw, self._ch = e.width, e.height
        self._redraw()

    def _redraw(self):
        w, h = self._cw, self._ch
        if w < 4 or h < 4:
            return
        self.delete("bg")
        self._bg_image = rounded_rect_photo(
            self, w, h, self._radius, self._app.c(self._fill)
        )
        self.create_image(0, 0, image=self._bg_image, anchor="nw", tags="bg")
        self.tag_lower("bg")
        self.itemconfigure(self._win, width=max(1, w - 2 * self._pad), height=max(1, h - 2 * self._pad))


class Btn(tk.Canvas):
    """Botão arredondado com hover."""
    VARIANTS = {
        "primary": ("accent", "accent_fg", "accent_hover"),
        "ghost": ("field", "fg", "line"),
        "danger": ("field", "danger", "line"),
        "nav": ("side", "side_fg", "side2"),
        "navon": ("side2", "side_fg", "side2"),
        "side": ("side2", "side_fg", "side_hover"),
    }

    def __init__(self, parent, app, text, command, variant="primary", on="card", size=10, bold=True,
                 height=38, padx=16, radius=12, anchor="center", expand=False, icon=None):
        self._app, self._text, self._cmd, self._variant, self._on = app, text, command, variant, on
        self._h, self._padx, self._radius, self._anchor, self._expand = height, padx, radius, anchor, expand
        self._icon, self._badge, self._hover = icon, None, False
        self._font = tkfont.Font(family=app.font, size=size, weight="bold" if bold else "normal")
        self._small = tkfont.Font(family=app.font, size=8, weight="bold")
        self._w0 = self._calc_width()
        super().__init__(parent, width=10 if expand else self._w0, height=height,
                         highlightthickness=0, bd=0, cursor="hand2")
        self.bind("<Configure>", lambda e: self.draw())
        self.bind("<Enter>", lambda e: self._set_hover(True))
        self.bind("<Leave>", lambda e: self._set_hover(False))
        self.bind("<ButtonRelease-1>", self._release)
        app.register(self)
        self.draw()

    def _calc_width(self):
        return self._font.measure(self._text) + 2 * self._padx + (30 if self._icon else 0)

    def _set_hover(self, value):
        self._hover = value
        self.draw()

    def _release(self, e):
        if 0 <= e.x <= self.winfo_width() and 0 <= e.y <= self._h and self._cmd:
            self._cmd()

    def set_text(self, text):
        self._text = text
        self._w0 = self._calc_width()
        if not self._expand:
            self.configure(width=self._w0)
        self.draw()

    def set_variant(self, variant):
        self._variant = variant
        self.draw()

    def set_badge(self, badge):
        self._badge = badge
        self.draw()

    def restyle(self):
        self.draw()

    def draw(self):
        a = self._app
        self.delete("all")
        w = self.winfo_width() if self._expand else self._w0
        if w < 8:
            return
        h = self._h
        bg_k, fg_k, hv_k = self.VARIANTS[self._variant]
        self.configure(bg=a.c(self._on))
        fill = a.c(hv_k if self._hover else bg_k)
        self._bg_image = rounded_rect_photo(self, w, h, self._radius, fill)
        self.create_image(w / 2, h / 2, image=self._bg_image, anchor="center")
        fg = a.c(fg_k)
        x = self._padx
        if self._anchor == "center":
            self.create_text(w / 2, h / 2, text=self._text, fill=fg, font=self._font)
        else:
            if self._icon:
                icon_color = a.c("accent") if self._variant in ("nav", "navon") else fg
                self.create_text(x + 2, h / 2, text=self._icon, anchor="w", fill=icon_color,
                                 font=(a.font, self._font.cget("size") + 1))
                x += 30
            self.create_text(x, h / 2, text=self._text, anchor="w", fill=fg, font=self._font)
        if self._badge:
            bw = self._small.measure(self._badge) + 16
            self.create_polygon(rounded_points(w - self._padx - bw, h / 2 - 10, w - self._padx, h / 2 + 10, 10),
                                smooth=True, fill=a.c("accent"), outline="")
            self.create_text(w - self._padx - bw / 2, h / 2, text=self._badge, fill=a.c("accent_fg"), font=self._small)


class Switch(tk.Canvas):
    def __init__(self, parent, app, value, command, on="side"):
        super().__init__(parent, width=42, height=24, highlightthickness=0, bd=0, cursor="hand2")
        self._app, self._value, self._cmd, self._on = app, bool(value), command, on
        self.bind("<Button-1>", self._toggle)
        app.register(self)
        self.restyle()

    def restyle(self):
        a = self._app
        self.configure(bg=a.c(self._on))
        self.delete("all")
        track = a.c("accent") if self._value else a.c("side_hover")
        self._switch_image = switch_photo(self, 42, 24, self._value, track)
        self.create_image(0, 0, image=self._switch_image, anchor="nw")

    def _toggle(self, _e=None):
        self._value = not self._value
        self.restyle()
        self._cmd(self._value)


class ScrollArea(tk.Frame):
    """Área rolável; coloque o conteúdo em `.body`."""

    def __init__(self, parent, app, bg="card"):
        super().__init__(parent, bd=0, highlightthickness=0)
        self._app, self._bgk = app, bg
        self.canvas = tk.Canvas(self, highlightthickness=0, bd=0, yscrollincrement=18)
        self.vsb = ttk.Scrollbar(self, orient="vertical", command=self.canvas.yview, style="App.Vertical.TScrollbar")
        self.canvas.configure(yscrollcommand=self.vsb.set)
        self.vsb.pack(side="right", fill="y")
        self.canvas.pack(side="left", fill="both", expand=True)
        self.body = tk.Frame(self.canvas, bd=0, highlightthickness=0)
        self._win = self.canvas.create_window(0, 0, window=self.body, anchor="nw")
        self.body.bind("<Configure>", lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        self.canvas.bind("<Configure>", lambda e: self.canvas.itemconfigure(self._win, width=e.width))
        app.register(self)
        self.restyle()

    def restyle(self):
        col = self._app.c(self._bgk)
        self.configure(bg=col)
        self.canvas.configure(bg=col)
        self.body.configure(bg=col)

    def scroll(self, units):
        lo, hi = self.canvas.yview()
        if hi - lo < 1.0:
            self.canvas.yview_scroll(units, "units")

    def clear(self):
        for w in self.body.winfo_children():
            w.destroy()
        self.canvas.yview_moveto(0)


class _IconRestyler:
    """Redesenha o ícone redondo do painel quando o tema muda."""

    def __init__(self, app, canvas, glyph):
        self.app, self.canvas, self.glyph = app, canvas, glyph

    def restyle(self):
        cv, a = self.canvas, self.app
        cv.delete("all")
        cv.create_oval(1, 1, 37, 37, fill=a.c("card"), outline="")
        cv.create_text(19, 19, text=self.glyph, fill=a.c("accent_text"), font=(a.font, 12))


# ----------------------------------------------------------------------------
# Aplicativo
# ----------------------------------------------------------------------------
class NotebookApp:
    def __init__(self, root):
        self.root = root
        self.config = read_config()
        self.root_dir = Path(self.config.get("root_dir", str(DEFAULT_ROOT))).expanduser()
        theme = self.config.get("theme", "light")
        self.theme_name = theme if theme in THEMES else "light"
        self.autosave_on = bool(self.config.get("autosave", True))
        self.font = self._pick_font()

        self.themed, self.themed_cfg, self.combos = [], [], []
        self.notes, self.struct_sig = [], None
        self.word_cache, self.counter_thread = {}, None
        self.current_file = None
        self.editing = self.dirty = self.meta_dirty = self.loading = False
        self.last_mtime, self.last_content = None, ""
        self.autosave_job = self.dirty_since = None
        self.view = "home"
        self._status = ("idle", "Pronto")
        self.row_refs = []
        self.foot_extra = ""
        self._search_job = self._words_job = None
        self._tick_n = 0
        self._ph = True
        self._search_guard = False

        root.title(APP_NAME)
        root.geometry("1280x800")
        root.minsize(1100, 720)
        self.style = ttk.Style()
        try:
            self.style.theme_use("clam")
        except tk.TclError:
            pass

        self.build_ui()
        self.apply_theme()
        self.update_sync_label()
        self.rescan()
        self.show_view("home")
        self.show_empty()

        root.bind_all("<MouseWheel>", self._on_wheel)
        root.bind_all("<Button-4>", self._on_wheel)
        root.bind_all("<Button-5>", self._on_wheel)
        for w in (root, self.body):
            w.bind("<Control-s>", self._key(self.save_now))
            w.bind("<Control-n>", self._key(self.new_note))
            w.bind("<Control-k>", self._key(self.focus_search))
        root.protocol("WM_DELETE_WINDOW", self.on_close)
        root.after(TICK_MS, self.tick)

    # ------------------------------------------------------------ helpers
    def _pick_font(self):
        try:
            fams = set(tkfont.families(self.root))
        except tk.TclError:
            fams = set()
        for f in ("Segoe UI", "Inter", "SF Pro Text", "Helvetica Neue", "DejaVu Sans", "Arial"):
            if f in fams:
                return f
        return "Helvetica"

    def c(self, key):
        return THEMES[self.theme_name][key]

    def register(self, widget):
        self.themed.append(widget)

    def T(self, widget, **opts):
        """Aplica cores da paleta e lembra delas para o troca de tema."""
        self.themed_cfg.append((widget, opts))
        widget.configure(**{k: self.c(v) for k, v in opts.items()})
        return widget

    def label(self, parent, text="", bg="card", fg="fg", size=10, bold=False, anchor="w", **kw):
        font = (self.font, size, "bold") if bold else (self.font, size)
        lb = tk.Label(parent, text=text, font=font, anchor=anchor, bd=0, **kw)
        return self.T(lb, bg=bg, fg=fg)

    def frame(self, parent, bg="card", **kw):
        return self.T(tk.Frame(parent, bd=0, highlightthickness=0, **kw), bg=bg)

    def make_combo(self, parent, var, width=18, readonly=True):
        cb = ttk.Combobox(parent, textvariable=var, width=width, style="App.TCombobox",
                          state="readonly" if readonly else "normal", font=(self.font, 10))
        self.combos.append(cb)
        return cb

    def bind_click(self, widget, fn):
        widget.bind("<Button-1>", lambda e: fn())
        try:
            widget.configure(cursor="hand2")
        except tk.TclError:
            pass
        for ch in widget.winfo_children():
            self.bind_click(ch, fn)

    def _key(self, fn):
        def handler(_e=None):
            fn()
            return "break"
        return handler

    def _on_wheel(self, event):
        try:
            w = self.root.winfo_containing(event.x_root, event.y_root)
        except (KeyError, tk.TclError):
            return
        while w is not None:
            if isinstance(w, ScrollArea):
                num = getattr(event, "num", 0)
                step = -1 if num == 4 else 1 if num == 5 else (-1 if event.delta > 0 else 1)
                w.scroll(step * 2)
                return
            w = getattr(w, "master", None)

    # ------------------------------------------------------------ tema
    def apply_styles(self):
        s, c = self.style, self.c
        s.configure("App.TCombobox", fieldbackground=c("field"), background=c("field"), foreground=c("fg"),
                    arrowcolor=c("muted"), bordercolor=c("line"), lightcolor=c("field"), darkcolor=c("field"),
                    selectbackground=c("field"), selectforeground=c("fg"), padding=6, relief="flat")
        s.map("App.TCombobox",
              fieldbackground=[("readonly", c("field")), ("focus", c("field"))],
              foreground=[("readonly", c("fg"))],
              selectbackground=[("readonly", c("field"))],
              selectforeground=[("readonly", c("fg"))],
              bordercolor=[("focus", c("accent"))])
        try:
            s.layout("App.Vertical.TScrollbar", [("Vertical.Scrollbar.trough", {
                "sticky": "ns", "children": [("Vertical.Scrollbar.thumb", {"expand": "1", "sticky": "nswe"})]})])
        except tk.TclError:
            pass
        s.configure("App.Vertical.TScrollbar", background=c("line"), troughcolor=c("card"),
                    bordercolor=c("card"), lightcolor=c("line"), darkcolor=c("line"), width=8, relief="flat")
        s.map("App.Vertical.TScrollbar", background=[("active", c("muted"))])

    def style_combo(self, cb):
        try:
            pop = cb.tk.eval(f"ttk::combobox::PopdownWindow {cb}")
            cb.tk.call(f"{pop}.f.l", "configure", "-background", self.c("card"), "-foreground", self.c("fg"),
                       "-selectbackground", self.c("select"), "-selectforeground", self.c("fg"),
                       "-borderwidth", 0, "-highlightthickness", 0)
        except tk.TclError:
            pass

    def apply_theme(self):
        self.root.configure(bg=self.c("bg"))
        self.apply_styles()
        for w, opts in list(self.themed_cfg):
            try:
                w.configure(**{k: self.c(v) for k, v in opts.items()})
            except tk.TclError:
                pass
        for w in list(self.themed):
            try:
                w.restyle()
            except tk.TclError:
                pass
        for cb in self.combos:
            self.style_combo(cb)
        self.draw_logo()
        self.draw_ws_icon()
        self.set_status(*self._status)
        self._paint_search_fg()
        self.theme_btn.set_text("☾" if self.theme_name == "light" else "☀")
        if self.view == "home":
            self.update_home()
        else:
            self.refresh_list()

    def toggle_theme(self):
        self.theme_name = "dark" if self.theme_name == "light" else "light"
        self.config["theme"] = self.theme_name
        write_config(self.config)
        self.apply_theme()

    # ------------------------------------------------------------ construção da interface
    def build_ui(self):
        self.root.grid_columnconfigure(1, weight=1)
        self.root.grid_rowconfigure(0, weight=1)
        self.title_var = tk.StringVar()
        self.edit_period_var = tk.StringVar()
        self.edit_course_var = tk.StringVar()
        self.period_var = tk.StringVar(value=ALL_PERIODS)
        self.course_var = tk.StringVar(value=ALL_COURSES)
        self.search_var = tk.StringVar()
        self.build_sidebar()
        self.build_main()
        for var, meta in ((self.title_var, True), (self.edit_period_var, True), (self.edit_course_var, True)):
            var.trace_add("write", lambda *_a, m=meta: self.touch(meta=m))
        self.edit_period_var.trace_add("write", lambda *_a: self.update_course_values())

    # -- barra lateral
    def build_sidebar(self):
        sb = tk.Frame(self.root, width=262, bd=0, highlightthickness=0)
        sb.grid(row=0, column=0, sticky="nsew")
        sb.pack_propagate(False)
        self.T(sb, bg="side")
        self.sidebar = sb

        row = self.frame(sb, "side")
        row.pack(fill="x", padx=22, pady=(26, 20))
        self.logo = self.T(tk.Canvas(row, width=30, height=26, highlightthickness=0, bd=0), bg="side")
        self.logo.pack(side="left")
        self.label(row, "Caderno", "side", "side_fg", 15, True).pack(side="left", padx=(10, 0))

        # pasta de trabalho
        ws = Surface(sb, self, fill="side2", on="side", radius=14, pad=10, height=64)
        ws.pack(fill="x", padx=16)
        self.ws_icon = self.T(tk.Canvas(ws.inner, width=36, height=36, highlightthickness=0, bd=0), bg="side2")
        self.ws_icon.pack(side="left")
        txt = self.frame(ws.inner, "side2")
        txt.pack(side="left", fill="x", expand=True, padx=(10, 0))
        self.label(txt, "Pasta de trabalho", "side2", "side_muted", 8).pack(anchor="w")
        self.ws_name = self.label(txt, "", "side2", "side_fg", 10, True)
        self.ws_name.pack(anchor="w")
        self.bind_click(ws, self.choose_root)

        # busca
        sr = Surface(sb, self, fill="side2", on="side", radius=12, pad=6, height=42)
        sr.pack(fill="x", padx=16, pady=(12, 0))
        self.label(sr.inner, "⌕", "side2", "side_muted", 13).pack(side="left", padx=(6, 4))
        self.search_entry = tk.Entry(sr.inner, textvariable=self.search_var, bd=0, relief="flat",
                                     highlightthickness=0, font=(self.font, 10))
        self.T(self.search_entry, bg="side2", insertbackground="side_fg")
        self.search_entry.pack(side="left", fill="both", expand=True)
        self.search_var.set("Buscar matérias…")
        self.search_entry.bind("<FocusIn>", self._search_focus_in)
        self.search_entry.bind("<FocusOut>", self._search_focus_out)
        self.search_var.trace_add("write", self._on_search)

        # navegação
        self.label(sb, "Navegação", "side", "side_muted", 9).pack(anchor="w", padx=22, pady=(24, 6))
        self.nav_home = Btn(sb, self, "Painel", lambda: self.show_view("home"), variant="navon", on="side",
                            size=10, bold=False, height=42, anchor="w", expand=True, icon="⌂")
        self.nav_home.pack(fill="x", padx=14, pady=2)
        self.nav_notes = Btn(sb, self, "Matérias", lambda: self.show_view("notes"), variant="nav", on="side",
                             size=10, bold=False, height=42, anchor="w", expand=True, icon="☰")
        self.nav_notes.pack(fill="x", padx=14, pady=2)

        # configurações (rodapé)
        bottom = self.frame(sb, "side")
        bottom.pack(side="bottom", fill="x", pady=(0, 18))
        self.label(bottom, "Configurações", "side", "side_muted", 9).pack(anchor="w", padx=22, pady=(0, 6))
        auto = self.frame(bottom, "side")
        auto.pack(fill="x", padx=22, pady=(2, 8))
        self.label(auto, "Salvamento automático", "side", "side_fg", 10).pack(side="left")
        Switch(auto, self, self.autosave_on, self.toggle_autosave, on="side").pack(side="right")
        Btn(bottom, self, "Alterar pasta do Drive…", self.choose_root, variant="nav", on="side", size=10,
            bold=False, height=40, anchor="w", expand=True, icon="⚙").pack(fill="x", padx=14, pady=2)
        Btn(bottom, self, "Abrir pasta no computador", self.open_root, variant="nav", on="side", size=10,
            bold=False, height=40, anchor="w", expand=True, icon="↗").pack(fill="x", padx=14, pady=2)

    def draw_logo(self):
        cv = self.logo
        cv.delete("all")
        for i, h in enumerate((14, 22, 18)):
            x = 2 + i * 9
            cv.create_polygon(rounded_points(x, 24 - h, x + 6, 24, 2), smooth=True, fill=self.c("accent"), outline="")

    def draw_ws_icon(self):
        cv = self.ws_icon
        cv.delete("all")
        cv.create_oval(1, 1, 35, 35, fill=self.c("accent"), outline="")
        cv.create_text(18, 18, text="☁", fill=self.c("accent_fg"), font=(self.font, 13))

    # -- busca (com texto de exemplo)
    def _search_focus_in(self, _e=None):
        if self._ph:
            self._search_guard = True
            self.search_var.set("")
            self._search_guard = False
            self._ph = False
            self._paint_search_fg()

    def _search_focus_out(self, _e=None):
        if not self.search_var.get().strip():
            self._search_guard = True
            self.search_var.set("Buscar matérias…")
            self._search_guard = False
            self._ph = True
            self._paint_search_fg()
            self.refresh_list()

    def _paint_search_fg(self):
        self.search_entry.configure(fg=self.c("side_muted" if self._ph else "side_fg"))

    def query(self):
        return "" if self._ph else self.search_var.get().strip().casefold()

    def _on_search(self, *_):
        if self._search_guard or self._ph:
            return
        if self._search_job:
            self.root.after_cancel(self._search_job)
        self._search_job = self.root.after(180, self._do_search)

    def _do_search(self):
        self._search_job = None
        if self.view != "notes":
            self.show_view("notes")
        else:
            self.refresh_list()

    def focus_search(self):
        self.show_view("notes")
        self.search_entry.focus_set()
        self.search_entry.selection_range(0, tk.END)

    # -- área principal
    def build_main(self):
        m = tk.Frame(self.root, bd=0, highlightthickness=0)
        m.grid(row=0, column=1, sticky="nsew")
        self.T(m, bg="bg")
        m.grid_columnconfigure(0, weight=1)
        m.grid_rowconfigure(1, weight=1)

        top = self.frame(m, "bg")
        top.grid(row=0, column=0, sticky="ew", padx=26, pady=(22, 14))
        self.title_lbl = self.label(top, "Painel", "bg", "fg", 21, True)
        self.title_lbl.pack(side="left")
        Btn(top, self, "＋  Nova matéria", self.new_note, variant="primary", on="bg", height=38, padx=18).pack(side="right")
        self.theme_btn = Btn(top, self, "☾", self.toggle_theme, variant="ghost", on="bg", size=12, height=38, padx=13)
        self.theme_btn.pack(side="right", padx=10)
        chip = Surface(top, self, fill="card", on="bg", radius=17, pad=8, width=196, height=36)
        chip.pack(side="right", padx=(0, 0))
        chip.pack_propagate(False)
        self.chip_dot = self.T(tk.Canvas(chip.inner, width=10, height=10, highlightthickness=0, bd=0), bg="card")
        self.chip_dot.pack(side="left", padx=(8, 8))
        self.chip_text = self.label(chip.inner, "Pronto", "card", "fg", 9)
        self.chip_text.pack(side="left")

        content = self.frame(m, "bg")
        content.grid(row=1, column=0, sticky="nsew", padx=26, pady=(0, 24))
        content.grid_columnconfigure(0, weight=1)
        content.grid_rowconfigure(0, weight=1)
        self.home = self.frame(content, "bg")
        self.home.grid(row=0, column=0, sticky="nsew")
        self.notes_view = self.frame(content, "bg")
        self.notes_view.grid(row=0, column=0, sticky="nsew")
        self.build_home()
        self.build_notes_view()

    def set_status(self, kind, text):
        self._status = (kind, text)
        key = {"saved": "accent", "dirty": "warn", "saving": "warn", "error": "danger", "idle": "muted"}[kind]
        self.chip_dot.delete("all")
        self.chip_dot.create_oval(1, 1, 9, 9, fill=self.c(key), outline="")
        self.chip_text.configure(text=text)

    # -- painel
    def build_home(self):
        h, gap = self.home, 14
        for i in range(3):
            h.grid_columnconfigure(i, weight=1, uniform="hc")
        h.grid_rowconfigure(2, weight=1)

        # cartões de resumo
        s1 = Surface(h, self, fill="dark", on="bg", radius=20, pad=18, height=118)
        s1.grid(row=0, column=0, sticky="nsew", padx=(0, gap), pady=(0, gap))
        left = self.frame(s1.inner, "dark")
        left.pack(side="left", fill="y")
        self.label(left, "Matérias salvas", "dark", "dark_muted", 9).pack(anchor="w")
        self.st_total = self.label(left, "0", "dark", "dark_fg", 26, True)
        self.st_total.pack(anchor="w", pady=(2, 0))
        self.st_total_sub = self.label(left, "", "dark", "accent", 9)
        self.st_total_sub.pack(anchor="w")
        self.mini1 = self.T(tk.Canvas(s1.inner, width=84, height=62, highlightthickness=0, bd=0), bg="dark")
        self.mini1.pack(side="right", anchor="s")

        s2 = Surface(h, self, fill="card", on="bg", radius=20, pad=18, height=118)
        s2.grid(row=0, column=1, sticky="nsew", padx=(0, gap), pady=(0, gap))
        left = self.frame(s2.inner, "card")
        left.pack(side="left", fill="y")
        self.label(left, "Disciplinas", "card", "muted", 9).pack(anchor="w")
        self.st_courses = self.label(left, "0", "card", "fg", 26, True)
        self.st_courses.pack(anchor="w", pady=(2, 0))
        self.st_courses_sub = self.label(left, "", "card", "muted", 9)
        self.st_courses_sub.pack(anchor="w")
        self.mini2 = self.T(tk.Canvas(s2.inner, width=84, height=62, highlightthickness=0, bd=0), bg="card")
        self.mini2.pack(side="right", anchor="s")

        s3 = Surface(h, self, fill="card", on="bg", radius=20, pad=18, height=118)
        s3.grid(row=0, column=2, sticky="nsew", pady=(0, gap))
        left = self.frame(s3.inner, "card")
        left.pack(side="left", fill="y")
        self.label(left, "Sincronização", "card", "muted", 9).pack(anchor="w")
        self.st_sync = self.label(left, "", "card", "fg", 17, True)
        self.st_sync.pack(anchor="w", pady=(8, 2))
        self.st_sync_sub = self.label(left, "", "card", "muted", 9)
        self.st_sync_sub.pack(anchor="w")
        self.sync_icon = self.T(tk.Canvas(s3.inner, width=44, height=44, highlightthickness=0, bd=0), bg="card")
        self.sync_icon.pack(side="right", anchor="n")

        # atividade
        cc = Surface(h, self, fill="card", on="bg", radius=20, pad=18, height=262)
        cc.grid(row=1, column=0, columnspan=2, sticky="nsew", padx=(0, gap), pady=(0, gap))
        head = self.frame(cc.inner, "card")
        head.pack(fill="x")
        col = self.frame(head, "card")
        col.pack(side="left")
        self.label(col, "Atividade de estudo", "card", "fg", 12, True).pack(anchor="w")
        self.label(col, "Matérias editadas por semana", "card", "muted", 9).pack(anchor="w")
        pill = self.label(head, "Últimas 8 semanas", "field", "fg", 9, padx=12, pady=5)
        pill.pack(side="right", anchor="n")
        self.chart_cv = self.T(tk.Canvas(cc.inner, highlightthickness=0, bd=0), bg="card")
        self.chart_cv.pack(fill="both", expand=True, pady=(10, 0))
        self.chart_cv.bind("<Configure>", lambda e: self.draw_chart())

        # palavras
        sc = Surface(h, self, fill="soft", on="bg", radius=20, pad=20, height=262)
        sc.grid(row=1, column=2, sticky="nsew", pady=(0, gap))
        self.st_words = self.label(sc.inner, "0", "soft", "soft_fg", 30, True)
        self.st_words.pack(anchor="w")
        self.st_words_sub = self.label(sc.inner, "palavras escritas", "soft", "accent_text", 10)
        self.st_words_sub.pack(anchor="w", pady=(0, 18))
        self.info1_t, self.info1_s = self._info_row(sc.inner, "✎")
        self.info2_t, self.info2_s = self._info_row(sc.inner, "☰")

        # recentes
        tc = Surface(h, self, fill="card", on="bg", radius=20, pad=18, height=250)
        tc.grid(row=2, column=0, columnspan=2, sticky="nsew", padx=(0, gap))
        head = self.frame(tc.inner, "card")
        head.pack(fill="x")
        self.label(head, "Matérias recentes", "card", "fg", 12, True).pack(side="left")
        Btn(head, self, "Ver todas", lambda: self.show_view("notes"), variant="ghost", on="card", size=9,
            height=30, padx=14, radius=10).pack(side="right")
        self.tbl = self.frame(tc.inner, "card")
        self.tbl.pack(fill="both", expand=True, pady=(8, 0))

        # por disciplina
        cc2 = Surface(h, self, fill="card", on="bg", radius=20, pad=18, height=250)
        cc2.grid(row=2, column=2, sticky="nsew")
        self.label(cc2.inner, "Por disciplina", "card", "fg", 12, True).pack(anchor="w")
        self.course_box_frame = self.frame(cc2.inner, "card")
        self.course_box_frame.pack(fill="both", expand=True, pady=(10, 0))

    def _info_row(self, parent, glyph):
        row = self.frame(parent, "soft")
        row.pack(fill="x", pady=7)
        ic = self.T(tk.Canvas(row, width=38, height=38, highlightthickness=0, bd=0), bg="soft")
        ic.pack(side="left")
        ic.create_oval(1, 1, 37, 37, fill=self.c("card"), outline="")
        ic.create_text(19, 19, text=glyph, fill=self.c("accent_text"), font=(self.font, 12))
        self.themed.append(_IconRestyler(self, ic, glyph))
        col = self.frame(row, "soft")
        col.pack(side="left", padx=(12, 0), fill="x", expand=True)
        t = self.label(col, "", "soft", "soft_fg", 10, True)
        t.pack(anchor="w")
        s = self.label(col, "", "soft", "accent_text", 9)
        s.pack(anchor="w")
        return t, s

    # -- visão de matérias
    def build_notes_view(self):
        v = self.notes_view
        v.grid_columnconfigure(0, minsize=330)
        v.grid_columnconfigure(1, weight=1)
        v.grid_rowconfigure(0, weight=1)

        lp = Surface(v, self, fill="card", on="bg", radius=20, pad=16, width=330, height=10)
        lp.grid(row=0, column=0, sticky="nsew", padx=(0, 14))
        head = self.frame(lp.inner, "card")
        head.pack(fill="x")
        self.label(head, "Matérias", "card", "fg", 14, True).pack(side="left")
        self.list_count = self.label(head, "", "card", "muted", 9)
        self.list_count.pack(side="right")
        self.label(lp.inner, "Período", "card", "muted", 9).pack(anchor="w", pady=(12, 3))
        self.period_box = self.make_combo(lp.inner, self.period_var, readonly=True)
        self.period_box.pack(fill="x")
        self.period_box.bind("<<ComboboxSelected>>", lambda e: self.on_filter_change())
        self.label(lp.inner, "Disciplina", "card", "muted", 9).pack(anchor="w", pady=(10, 3))
        self.course_box = self.make_combo(lp.inner, self.course_var, readonly=True)
        self.course_box.pack(fill="x")
        self.course_box.bind("<<ComboboxSelected>>", lambda e: self.refresh_list())
        self.list_hint = self.label(lp.inner, "", "card", "muted", 9)
        self.list_hint.pack(anchor="w", pady=(10, 4))
        self.scroll = ScrollArea(lp.inner, self, "card")
        self.scroll.pack(fill="both", expand=True)

        ed = Surface(v, self, fill="card", on="bg", radius=20, pad=22)
        ed.grid(row=0, column=1, sticky="nsew")
        ed.inner.grid_columnconfigure(0, weight=1)
        ed.inner.grid_rowconfigure(0, weight=1)

        # estado vazio
        self.empty_frame = self.frame(ed.inner, "card")
        self.empty_frame.grid(row=0, column=0, sticky="nsew")
        box = self.frame(self.empty_frame, "card")
        box.place(relx=0.5, rely=0.5, anchor="center")
        self.label(box, "✎", "card", "accent_text", 34, anchor="center").pack()
        self.label(box, "Nenhuma matéria aberta", "card", "fg", 15, True, anchor="center").pack(pady=(6, 2))
        self.label(box, "Escolha uma matéria na lista ou crie uma nova.", "card", "muted", 10, anchor="center").pack()
        Btn(box, self, "＋  Nova matéria", self.new_note, variant="primary", on="card", height=40, padx=20).pack(pady=(16, 0))

        # editor
        self.editor_frame = self.frame(ed.inner, "card")
        self.editor_frame.grid(row=0, column=0, sticky="nsew")
        f = self.editor_frame
        f.grid_columnconfigure(0, weight=1)
        f.grid_rowconfigure(4, weight=1)
        self.title_entry = tk.Entry(f, textvariable=self.title_var, bd=0, relief="flat", highlightthickness=0,
                                    font=(self.font, 21, "bold"))
        self.T(self.title_entry, bg="card", fg="fg", insertbackground="fg")
        self.title_entry.grid(row=0, column=0, sticky="ew")
        self.title_entry.bind("<Return>", lambda e: (self.save_soon(), self.body.focus_set()))
        self.title_entry.bind("<FocusOut>", lambda e: self.save_soon())
        self.location_lbl = self.label(f, "", "card", "muted", 9)
        self.location_lbl.grid(row=1, column=0, sticky="w", pady=(2, 0))

        meta = self.frame(f, "card")
        meta.grid(row=2, column=0, sticky="w", pady=(12, 0))
        self.label(meta, "Período", "card", "muted", 9).pack(side="left")
        self.edit_period_box = self.make_combo(meta, self.edit_period_var, width=16, readonly=False)
        self.edit_period_box.pack(side="left", padx=(8, 18))
        self.label(meta, "Disciplina", "card", "muted", 9).pack(side="left")
        self.edit_course_box = self.make_combo(meta, self.edit_course_var, width=26, readonly=False)
        self.edit_course_box.pack(side="left", padx=(8, 0))
        for cb in (self.edit_period_box, self.edit_course_box):
            cb.bind("<FocusOut>", lambda e: self.save_soon())
            cb.bind("<<ComboboxSelected>>", lambda e: self.save_soon())

        sep = self.frame(f, "line", height=1)
        sep.grid(row=3, column=0, sticky="ew", pady=(14, 0))

        wrap = self.frame(f, "card")
        wrap.grid(row=4, column=0, sticky="nsew", pady=(6, 0))
        self.body = tk.Text(wrap, wrap="word", undo=True, bd=0, relief="flat", highlightthickness=0, padx=2,
                            pady=12, font=(self.font, 12), spacing1=2, spacing2=3, spacing3=6)
        self.T(self.body, bg="card", fg="fg", insertbackground="fg", selectbackground="select",
               selectforeground="fg", inactiveselectbackground="select")
        bs = ttk.Scrollbar(wrap, orient="vertical", command=self.body.yview, style="App.Vertical.TScrollbar")
        self.body.configure(yscrollcommand=bs.set)
        bs.pack(side="right", fill="y")
        self.body.pack(side="left", fill="both", expand=True)
        self.body.bind("<<Modified>>", self.on_body_modified)

        foot = self.frame(f, "card")
        foot.grid(row=5, column=0, sticky="ew", pady=(8, 0))
        self.foot_lbl = self.label(foot, "", "card", "muted", 9)
        self.foot_lbl.pack(side="left")
        Btn(foot, self, "Salvar agora", self.save_now, variant="primary", on="card", height=36, padx=16).pack(side="right")
        Btn(foot, self, "Abrir no Word", self.open_in_word, variant="ghost", on="card", height=36, padx=14).pack(side="right", padx=8)
        Btn(foot, self, "Excluir", self.delete_note, variant="danger", on="card", height=36, padx=14).pack(side="right")

    def show_empty(self):
        self.editor_frame.grid_remove()
        self.empty_frame.grid()
        self.empty_frame.tkraise()

    def show_editor(self):
        self.empty_frame.grid_remove()
        self.editor_frame.grid()
        self.editor_frame.tkraise()

    def show_view(self, name):
        self.view = name
        if name == "home":
            self.notes_view.grid_remove()
            self.home.grid()
            self.title_lbl.configure(text="Painel")
            self.update_home()
        else:
            self.home.grid_remove()
            self.notes_view.grid()
            self.title_lbl.configure(text="Matérias")
            self.refresh_list()
        self.nav_home.set_variant("navon" if name == "home" else "nav")
        self.nav_notes.set_variant("navon" if name == "notes" else "nav")


    # ------------------------------------------------------------ pasta e arquivos
    def update_sync_label(self):
        name = self.root_dir.name or str(self.root_dir)
        self.ws_name.configure(text=trunc(name, 20))

    def choose_root(self):
        if not self.leave_ok():
            return
        chosen = filedialog.askdirectory(title="Escolha a pasta sincronizada do Google Drive")
        if not chosen:
            return
        self.root_dir = Path(chosen)
        self.config["root_dir"] = str(self.root_dir)
        write_config(self.config)
        self.cancel_autosave()
        self.editing = self.dirty = self.meta_dirty = False
        self.current_file = None
        self.show_empty()
        self.update_sync_label()
        self.set_status("idle", "Pasta atualizada")
        self.rescan()
        self.highlight_current()

    def open_root(self):
        try:
            self.root_dir.mkdir(parents=True, exist_ok=True)
            self._open_path(self.root_dir)
        except Exception as e:
            messagebox.showerror("Não foi possível abrir", str(e))

    def open_in_word(self):
        if self.current_file is None or not self.current_file.exists():
            messagebox.showinfo("Abrir no Word", "Aguarde o salvamento da matéria para abri-la no Word.")
            return
        if self.dirty and not self.persist():
            return
        try:
            self._open_path(self.current_file)
        except Exception as e:
            messagebox.showerror("Não foi possível abrir", str(e))

    @staticmethod
    def _open_path(path):
        if sys.platform.startswith("win"):
            os.startfile(str(path))
        elif sys.platform == "darwin":
            subprocess.Popen(["open", str(path)])
        else:
            subprocess.Popen(["xdg-open", str(path)])

    def rescan(self):
        """Relê a pasta. Só reconstrói a lista quando arquivos mudam."""
        notes = []
        root = self.root_dir
        if root.exists():
            try:
                for p in root.rglob("*.docx"):
                    if p.name.startswith(("~$", ".")):
                        continue
                    try:
                        rel = p.relative_to(root)
                        st = p.stat()
                    except (ValueError, OSError):
                        continue
                    if len(rel.parts) < 3:
                        continue
                    notes.append({"path": p, "title": p.stem, "period": rel.parts[0],
                                  "course": rel.parts[1], "mtime": st.st_mtime})
            except OSError:
                pass
        notes.sort(key=lambda n: n["mtime"], reverse=True)
        sig = tuple(sorted(str(n["path"]) for n in notes))
        structure_changed = sig != self.struct_sig
        self.notes, self.struct_sig = notes, sig
        self.nav_notes.set_badge(str(len(notes)) if notes else None)
        self.refresh_filters()
        if self.view == "home":
            self.update_home()
        elif structure_changed:
            self.refresh_list()
        self.highlight_current()
        self.start_counter()

    def refresh_filters(self):
        try:
            periods = sorted((p.name for p in self.root_dir.iterdir()
                              if p.is_dir() and not p.name.startswith(".")), key=str.casefold)
        except OSError:
            periods = []
        self.period_box["values"] = [ALL_PERIODS] + periods
        if self.period_var.get() not in self.period_box["values"]:
            self.period_var.set(ALL_PERIODS)
        sel = self.period_var.get()
        courses = set()
        for p in (periods if sel == ALL_PERIODS else [sel]):
            try:
                courses.update(d.name for d in (self.root_dir / p).iterdir() if d.is_dir())
            except OSError:
                pass
        self.course_box["values"] = [ALL_COURSES] + sorted(courses, key=str.casefold)
        if self.course_var.get() not in self.course_box["values"]:
            self.course_var.set(ALL_COURSES)
        self.edit_period_box["values"] = periods
        self.update_course_values()

    def update_course_values(self):
        period = clean(self.edit_period_var.get(), "")
        values = []
        if period:
            try:
                values = sorted((d.name for d in (self.root_dir / period).iterdir() if d.is_dir()), key=str.casefold)
            except OSError:
                values = []
        self.edit_course_box["values"] = values

    def on_filter_change(self):
        self.course_var.set(ALL_COURSES)
        self.refresh_filters()
        self.refresh_list()

    # ------------------------------------------------------------ lista de matérias
    def refresh_list(self):
        self.scroll.clear()
        self.row_refs = []
        period, course, q = self.period_var.get(), self.course_var.get(), self.query()
        shown = []
        for n in self.notes:
            if period != ALL_PERIODS and n["period"] != period:
                continue
            if course != ALL_COURSES and n["course"] != course:
                continue
            if q and q not in n["title"].casefold() and q not in n["course"].casefold() \
                    and q not in n["period"].casefold():
                continue
            shown.append(n)
        total = len(shown)
        self.list_count.configure(text=f"{total} {'matéria' if total == 1 else 'matérias'}")
        self.list_hint.configure(text=f"Resultados para “{self.search_var.get().strip()}”" if q else "Mais recentes primeiro")
        if not shown:
            msg = "Nenhuma matéria encontrada." if (q or period != ALL_PERIODS or course != ALL_COURSES) \
                else "Ainda não há matérias.\nClique em “Nova matéria” para começar."
            tk.Label(self.scroll.body, text=msg, font=(self.font, 10), justify="left", anchor="w",
                     bg=self.c("card"), fg=self.c("muted")).pack(fill="x", pady=16, padx=4)
            return
        for n in shown[:250]:
            self.build_list_row(n)
        if total > 250:
            tk.Label(self.scroll.body, text=f"Mostrando 250 de {total}. Use a busca para refinar.",
                     font=(self.font, 9), bg=self.c("card"), fg=self.c("muted")).pack(pady=8)

    def build_list_row(self, n):
        f = self.font
        selected = self.current_file is not None and same_path(n["path"], self.current_file)
        bg = self.c("soft") if selected else self.c("card")
        row = tk.Frame(self.scroll.body, bg=bg, cursor="hand2")
        row.pack(fill="x", pady=2)
        inner = tk.Frame(row, bg=bg)
        inner.pack(fill="x", padx=12, pady=9)
        t = tk.Label(inner, text=trunc(n["title"], 32), font=(f, 10, "bold"), bg=bg, fg=self.c("fg"), anchor="w")
        t.pack(fill="x")
        sub = tk.Label(inner, text=f"{trunc(n['course'], 22)} · {n['period']} · {rel_time(n['mtime'])}",
                       font=(f, 8), bg=bg, fg=self.c("muted"), anchor="w")
        sub.pack(fill="x")
        ref = {"path": n["path"], "widgets": [row, inner, t, sub], "normal": bg}
        self.row_refs.append(ref)
        path = n["path"]
        self.bind_click(row, lambda p=path: self.open_note(p))
        for w in ref["widgets"]:
            w.bind("<Enter>", lambda e, r=ref: self._paint_row(r, self.c("field") if r["normal"] == self.c("card") else r["normal"]))
            w.bind("<Leave>", lambda e, r=ref: self._paint_row(r, r["normal"]))

    @staticmethod
    def _paint_row(ref, color):
        for w in ref["widgets"]:
            try:
                w.configure(bg=color)
            except tk.TclError:
                pass

    def highlight_current(self):
        for ref in self.row_refs:
            sel = self.current_file is not None and same_path(ref["path"], self.current_file)
            ref["normal"] = self.c("soft") if sel else self.c("card")
            self._paint_row(ref, ref["normal"])

    # ------------------------------------------------------------ painel: números e gráficos
    def weekly_counts(self, n=8):
        today = datetime.now().date()
        counts = [0] * n
        for x in self.notes:
            d = max(0, (today - datetime.fromtimestamp(x["mtime"]).date()).days)
            if d // 7 < n:
                counts[n - 1 - d // 7] += 1
        labels = [(today - timedelta(days=7 * (n - 1 - i) + 6)).strftime("%d/%m") for i in range(n)]
        return list(zip(labels, counts))

    def draw_mini(self, cv, values, bg_key, color_key, hi_key):
        cv.delete("all")
        cv.configure(bg=self.c(bg_key))
        w, h = int(float(cv.cget("width"))), int(float(cv.cget("height")))
        n, mx, gap = len(values), max(max(values), 1), 5
        if n == 0:
            return
        bw = (w - gap * (n - 1)) / n
        for i, v in enumerate(values):
            bh = max(5, (h - 4) * v / mx)
            x0 = i * (bw + gap)
            color = self.c(hi_key) if i == n - 1 else self.c(color_key)
            cv.create_polygon(rounded_points(x0, h - bh, x0 + bw, h, 3), smooth=True, fill=color, outline="")

    def update_home(self):
        notes = self.notes
        today = datetime.now().date()
        days, week, today_cnt = [0] * 7, 0, 0
        per_course, periods = {}, set()
        for x in notes:
            d = (today - datetime.fromtimestamp(x["mtime"]).date()).days
            if 0 <= d < 7:
                days[6 - d] += 1
                week += 1
            if d <= 0:
                today_cnt += 1
            per_course[x["course"]] = per_course.get(x["course"], 0) + 1
            periods.add(x["period"])
        self.st_total.configure(text=str(len(notes)))
        self.st_total_sub.configure(text=f"▲ {week} editadas nos últimos 7 dias" if week else "Nada editado esta semana")
        self.draw_mini(self.mini1, days, "dark", "accent", "accent")
        self.st_courses.configure(text=str(len(per_course)))
        self.st_courses_sub.configure(text=f"em {len(periods)} {'período' if len(periods) == 1 else 'períodos'}")
        top = sorted(per_course.values(), reverse=True)[:6] or [0]
        self.draw_mini(self.mini2, top[::-1], "card", "bar", "accent")

        drive = is_drive_path(self.root_dir)
        self.st_sync.configure(text="Google Drive" if drive else "Pasta local")
        auto = "Salvamento automático ativo" if self.autosave_on else "Salvamento manual"
        self.st_sync_sub.configure(text=auto if drive else "Escolha uma pasta do Drive")
        cv = self.sync_icon
        cv.delete("all")
        cv.create_oval(2, 2, 42, 42, fill=self.c("accent" if drive else "field"), outline="")
        cv.create_text(22, 22, text="✓" if drive else "!", fill=self.c("accent_fg" if drive else "muted"),
                       font=(self.font, 15, "bold"))

        if notes:
            newest = notes[0]
            self.info1_t.configure(text=trunc(newest["title"], 26))
            self.info1_s.configure(text=f"Última edição · {rel_time(newest['mtime'])}")
            name, cnt = max(per_course.items(), key=lambda kv: kv[1])
            self.info2_t.configure(text=trunc(name, 26))
            self.info2_s.configure(text=f"Disciplina com mais matérias ({cnt})")
        else:
            self.info1_t.configure(text="Nenhuma edição ainda")
            self.info1_s.configure(text="Crie sua primeira matéria")
            self.info2_t.configure(text="Sem disciplinas")
            self.info2_s.configure(text="Elas aparecem aqui")
        self.update_words_stat()
        self.fill_table()
        self.fill_courses(per_course)
        self.draw_chart()

    def fill_table(self):
        for w in self.tbl.winfo_children():
            w.destroy()
        cols = (("Matéria", 30), ("Disciplina", 19), ("Período", 12), ("Editada", 13))
        header = tk.Frame(self.tbl, bg=self.c("card"))
        header.pack(fill="x")
        for i, (name, width) in enumerate(cols):
            tk.Label(header, text=name, width=width, anchor="w", font=(self.font, 9), bg=self.c("card"),
                     fg=self.c("muted")).grid(row=0, column=i, sticky="w")
        if not self.notes:
            tk.Label(self.tbl, text="Nenhuma matéria ainda. Use “Nova matéria” para criar a primeira.",
                     font=(self.font, 10), bg=self.c("card"), fg=self.c("muted")).pack(anchor="w", pady=20)
            return
        for n in self.notes[:5]:
            row = tk.Frame(self.tbl, bg=self.c("card"), cursor="hand2")
            row.pack(fill="x", pady=1)
            cells = (trunc(n["title"], 30), trunc(n["course"], 19), trunc(n["period"], 12), rel_time(n["mtime"]))
            labels = []
            for i, ((name, width), text) in enumerate(zip(cols, cells)):
                lb = tk.Label(row, text=text, width=width, anchor="w", pady=7, bg=self.c("card"),
                              fg=self.c("fg") if i == 0 else self.c("muted"),
                              font=(self.font, 10, "bold") if i == 0 else (self.font, 10))
                lb.grid(row=0, column=i, sticky="w")
                labels.append(lb)
            ref = {"widgets": [row] + labels}
            self.bind_click(row, lambda p=n["path"]: self.open_note(p))
            for w in ref["widgets"]:
                w.bind("<Enter>", lambda e, r=ref: self._paint_row(r, self.c("field")))
                w.bind("<Leave>", lambda e, r=ref: self._paint_row(r, self.c("card")))

    def fill_courses(self, per_course):
        box = self.course_box_frame
        for w in box.winfo_children():
            w.destroy()
        items = sorted(per_course.items(), key=lambda kv: (-kv[1], kv[0].casefold()))[:5]
        if not items:
            tk.Label(box, text="As disciplinas aparecem aqui.", font=(self.font, 10), bg=self.c("card"),
                     fg=self.c("muted")).pack(anchor="w", pady=14)
            return
        mx = items[0][1]
        for name, cnt in items:
            line = tk.Frame(box, bg=self.c("card"))
            line.pack(fill="x", pady=(0, 9))
            top = tk.Frame(line, bg=self.c("card"))
            top.pack(fill="x")
            tk.Label(top, text=trunc(name, 24), font=(self.font, 10), bg=self.c("card"), fg=self.c("fg"),
                     anchor="w").pack(side="left")
            tk.Label(top, text=str(cnt), font=(self.font, 10, "bold"), bg=self.c("card"),
                     fg=self.c("muted")).pack(side="right")
            track = tk.Frame(line, bg=self.c("line"), height=6)
            track.pack(fill="x", pady=(3, 0))
            track.pack_propagate(False)
            tk.Frame(track, bg=self.c("accent")).place(x=0, y=0, relheight=1, relwidth=max(cnt / mx, 0.04))

    def draw_chart(self):
        cv = self.chart_cv
        cv.delete("all")
        cv.configure(bg=self.c("card"))
        w, h = cv.winfo_width(), cv.winfo_height()
        if w < 120 or h < 90:
            return
        data = self.weekly_counts(8)
        ymax = max(4, ((max(v for _, v in data) + 3) // 4) * 4)
        left, right, top, bottom = 32, w - 6, 28, h - 24
        for i in range(5):
            y = bottom - (bottom - top) * i / 4
            cv.create_line(left, y, right, y, fill=self.c("line"))
            cv.create_text(left - 8, y, text=str(int(ymax * i / 4)), anchor="e", fill=self.c("muted"), font=(self.font, 8))
        n = len(data)
        slot = (right - left) / n
        bw = min(30, slot * 0.5)
        bubble = None
        for i, (lab, v) in enumerate(data):
            cx = left + slot * (i + 0.5)
            last = i == n - 1
            bh = (bottom - top) * v / ymax if v else 3
            y0 = bottom - bh
            color = self.c("dark") if (last and v) else self.c("accent") if v else self.c("bar")
            cv.create_polygon(rounded_points(cx - bw / 2, y0, cx + bw / 2, bottom, 5), smooth=True, fill=color, outline="")
            cv.create_text(cx, bottom + 12, text="esta sem." if last else lab, fill=self.c("muted"), font=(self.font, 8))
            if last and v:
                bubble = (cx, y0, v)
        if bubble:
            cx, y0, v = bubble
            cv.create_line(left, y0, right, y0, fill=self.c("muted"), dash=(3, 3))
            text = f"{v} {'matéria' if v == 1 else 'matérias'}"
            tw = tkfont.Font(family=self.font, size=8, weight="bold").measure(text) + 18
            bx0 = min(max(cx - tw / 2, left), right - tw)
            cv.create_polygon(rounded_points(bx0, y0 - 28, bx0 + tw, y0 - 8, 8), smooth=True, fill=self.c("dark"), outline="")
            cv.create_text(bx0 + tw / 2, y0 - 18, text=text, fill=self.c("dark_fg"), font=(self.font, 8, "bold"))

    # -- contagem de palavras em segundo plano (para não travar a janela)
    def start_counter(self):
        if self.counter_thread is not None and self.counter_thread.is_alive():
            return
        todo = [(n["path"], n["mtime"]) for n in self.notes
                if (self.word_cache.get(str(n["path"])) or (None, 0))[0] != n["mtime"]]
        if not todo:
            return

        def work():
            for path, mtime in todo:
                try:
                    count = count_words(read_note(path))
                except Exception:
                    count = 0
                self.word_cache[str(path)] = (mtime, count)

        self.counter_thread = threading.Thread(target=work, daemon=True)
        self.counter_thread.start()
        self.root.after(500, self._poll_counter)

    def _poll_counter(self):
        self.update_words_stat()
        if self.counter_thread is not None and self.counter_thread.is_alive():
            self.root.after(500, self._poll_counter)

    def update_words_stat(self):
        total = 0
        for x in self.notes:
            entry = self.word_cache.get(str(x["path"]))
            if entry and entry[0] == x["mtime"]:
                total += entry[1]
        counting = self.counter_thread is not None and self.counter_thread.is_alive()
        self.st_words.configure(text=f"{total:,}".replace(",", "."))
        self.st_words_sub.configure(text="palavras escritas" + (" (contando…)" if counting else ""))

    # ------------------------------------------------------------ edição e salvamento automático
    def on_body_modified(self, _e=None):
        if self.body.edit_modified():
            self.body.edit_modified(False)
            self.touch()

    def touch(self, meta=False):
        """Chamado a cada alteração do usuário."""
        if self.loading or not self.editing:
            return
        self.dirty = True
        if meta:
            self.meta_dirty = True
        if self.dirty_since is None:
            self.dirty_since = time.time()
        self.update_footer_soon()
        if self.autosave_on:
            self.set_status("dirty", "Editando…")
            self.schedule_autosave()
        else:
            self.set_status("dirty", "Alterações não salvas")

    def cancel_autosave(self):
        if self.autosave_job is not None:
            try:
                self.root.after_cancel(self.autosave_job)
            except tk.TclError:
                pass
            self.autosave_job = None

    def schedule_autosave(self, delay=None):
        self.cancel_autosave()
        if delay is None:
            delay = AUTOSAVE_META_MS if self.meta_dirty else AUTOSAVE_BODY_MS
            waited = (time.time() - (self.dirty_since or time.time())) * 1000
            delay = int(max(200, min(delay, AUTOSAVE_MAX_WAIT_MS - waited)))
        self.autosave_job = self.root.after(delay, self.autosave_tick)

    def save_soon(self):
        if self.editing and self.dirty and self.autosave_on:
            self.schedule_autosave(250)

    def autosave_tick(self):
        self.autosave_job = None
        if self.editing and self.dirty and not self.persist():
            self.autosave_job = self.root.after(6000, self.autosave_tick)

    def save_now(self):
        if not self.editing:
            return
        self.cancel_autosave()
        self.persist()

    def toggle_autosave(self, value):
        self.autosave_on = bool(value)
        self.config["autosave"] = self.autosave_on
        write_config(self.config)
        if self.editing and self.dirty:
            if self.autosave_on:
                self.set_status("dirty", "Editando…")
                self.schedule_autosave()
            else:
                self.cancel_autosave()
                self.set_status("dirty", "Alterações não salvas")
        if self.view == "home":
            self.update_home()

    def update_footer_soon(self):
        if self._words_job is None:
            self._words_job = self.root.after(250, self.update_footer)

    def update_footer(self):
        self._words_job = None
        if not self.editing:
            return
        text = self.body.get("1.0", "end-1c")
        s = f"{count_words(text)} palavras · {len(text)} caracteres"
        if self.foot_extra:
            s += f"  ·  {self.foot_extra}"
        self.foot_lbl.configure(text=s)

    def persist(self):
        """Grava a matéria aberta no DOCX. Devolve True se deu certo."""
        if not self.editing:
            return True
        raw_title = self.title_var.get().strip()
        body = self.body.get("1.0", "end-1c")
        if self.current_file is None and not raw_title and not body.strip():
            self.dirty = self.meta_dirty = False
            self.dirty_since = None
            self.set_status("idle", "Nova matéria — comece a escrever")
            return True

        period = clean(self.edit_period_var.get(), "") or clean(self.config.get("last_period", ""), "1º Período")
        course = clean(self.edit_course_var.get(), "Geral")
        title = clean(raw_title or first_line(body), "Sem título")
        folder = self.root_dir / period / course
        target = folder / f"{title}.docx"
        old = self.current_file
        same = old is not None and same_path(old, target)
        if not same and target.exists():          # nunca sobrescreve outra matéria
            target = unique_path(target)
            title = target.stem

        self.set_status("saving", "Salvando…")
        self.root.update_idletasks()
        self.foot_extra = ""
        try:
            folder.mkdir(parents=True, exist_ok=True)
            if same and target.exists() and self.last_mtime is not None:
                try:
                    if target.stat().st_mtime != self.last_mtime:
                        disk = read_note(target)
                        if disk != self.last_content and disk != body:
                            backup = unique_path(target.with_name(f"{target.stem} (versão do Drive).docx"))
                            shutil.copy2(str(target), str(backup))
                            self.foot_extra = "o Drive tinha outra versão; guardei uma cópia"
                except Exception:
                    pass
            write_docx(target, title, period, course, body)
            if old is not None and not same and old.exists():
                try:
                    old.unlink()
                except OSError:
                    pass
            mtime = target.stat().st_mtime
            try:
                saved_content = read_note(target)
            except Exception:
                saved_content = body
        except PermissionError:
            self.set_status("error", "Arquivo em uso. Tentando de novo…")
            return False
        except Exception as e:
            self.set_status("error", "Erro ao salvar")
            self.foot_extra = trunc(str(e), 70)
            self.update_footer()
            return False

        self.current_file, self.last_mtime, self.last_content = target, mtime, saved_content
        self.dirty = self.meta_dirty = False
        self.dirty_since = None
        self.loading = True
        try:
            if self.title_var.get() != title:
                self.title_var.set(title)
            if self.edit_period_var.get() != period:
                self.edit_period_var.set(period)
            if self.edit_course_var.get() != course:
                self.edit_course_var.set(course)
        finally:
            self.loading = False
        if self.config.get("last_period") != period:
            self.config["last_period"] = period
            write_config(self.config)
        try:
            self.location_lbl.configure(text=str(target.relative_to(self.root_dir)))
        except ValueError:
            self.location_lbl.configure(text=str(target))
        self.set_status("saved", f"Salvo às {datetime.now():%H:%M}")
        self.update_footer()
        self.rescan()
        return True

    def leave_ok(self):
        """Chamado antes de trocar de matéria ou fechar. True = pode continuar."""
        if not (self.editing and self.dirty):
            return True
        self.cancel_autosave()
        if not self.autosave_on:
            answer = messagebox.askyesnocancel("Alterações não salvas", "Deseja salvar as alterações antes de continuar?")
            if answer is None:
                return False
            if answer is False:
                self.dirty = self.meta_dirty = False
                return True
        if self.persist():
            return True
        return messagebox.askyesno(
            "Não foi possível salvar",
            "Não consegui salvar esta matéria (ela pode estar aberta no Word).\n\nContinuar mesmo assim e perder as alterações?")

    # ------------------------------------------------------------ abrir / criar / excluir
    def open_note(self, path):
        if self.current_file is not None and same_path(path, self.current_file) and self.editing:
            self.show_view("notes")
            return
        if not self.leave_ok():
            return
        try:
            content = read_note(path)
            mtime = path.stat().st_mtime
        except Exception as e:
            messagebox.showerror("Erro ao abrir DOCX", str(e))
            return
        try:
            rel = path.relative_to(self.root_dir)
            period, course = rel.parts[0], rel.parts[1]
        except (ValueError, IndexError):
            period, course = "", ""
        self._load_into_editor(path.stem, period, course, content)
        self.current_file, self.last_mtime, self.last_content = path, mtime, content
        try:
            self.location_lbl.configure(text=str(path.relative_to(self.root_dir)))
        except ValueError:
            self.location_lbl.configure(text=str(path))
        self.editing = True
        self.dirty = self.meta_dirty = False
        self.dirty_since = None
        self.foot_extra = ""
        self.show_editor()
        self.show_view("notes")
        self.update_footer()
        self.set_status("saved", "Matéria carregada")

    def _load_into_editor(self, title, period, course, content):
        self.loading = True
        try:
            self.title_var.set(title)
            self.edit_period_var.set(period)
            self.edit_course_var.set(course)
            self.body.delete("1.0", tk.END)
            self.body.insert("1.0", content)
            self.body.edit_modified(False)
            self.body.edit_reset()
            self.body.mark_set("insert", "1.0")
            self.body.yview_moveto(0)
        finally:
            self.loading = False

    def new_note(self):
        if not self.leave_ok():
            return
        period = self.period_var.get()
        if period in ("", ALL_PERIODS):
            period = clean(self.config.get("last_period", ""), "1º Período")
        course = self.course_var.get()
        if course in ("", ALL_COURSES):
            course = ""
        self._load_into_editor("", period, course, "")
        self.current_file, self.last_mtime, self.last_content = None, None, ""
        self.editing = True
        self.dirty = self.meta_dirty = False
        self.dirty_since = None
        self.foot_extra = ""
        self.cancel_autosave()
        self.location_lbl.configure(text="Nova matéria — o arquivo é criado assim que você começa a escrever.")
        self.show_editor()
        self.show_view("notes")
        self.highlight_current()
        self.update_footer()
        self.set_status("idle", "Nova matéria")
        self.title_entry.focus_set()

    def delete_note(self):
        if self.current_file is None or not self.current_file.exists():
            messagebox.showinfo("Excluir matéria", "Esta matéria ainda não foi salva.")
            return
        if not messagebox.askyesno("Confirmar exclusão", f"Excluir permanentemente este arquivo?\n\n{self.current_file}"):
            return
        try:
            path = self.current_file
            self.cancel_autosave()
            path.unlink()
        except OSError as e:
            messagebox.showerror("Erro ao excluir", str(e))
            return
        self.current_file = None
        self.editing = self.dirty = self.meta_dirty = False
        self.show_empty()
        self.set_status("idle", f"Excluída: {trunc(path.stem, 18)}")
        self.rescan()
        self.refresh_list()

    # ------------------------------------------------------------ rotinas periódicas
    def check_external(self):
        """Se o Drive trouxe uma versão nova do arquivo aberto e não há edição pendente, recarrega."""
        p = self.current_file
        if not (self.editing and p is not None):
            return
        try:
            m = p.stat().st_mtime
        except OSError:
            return
        if self.last_mtime is None or m == self.last_mtime:
            return
        try:
            content = read_note(p)
        except Exception:
            return                      # arquivo ainda sendo escrito pelo Drive
        if content == self.last_content:
            self.last_mtime = m
            return
        if self.dirty:
            return                      # o conflito é tratado no próximo salvamento
        cursor = self.body.index("insert")
        self._load_into_editor(self.title_var.get(), self.edit_period_var.get(), self.edit_course_var.get(), content)
        try:
            self.body.mark_set("insert", cursor)
        except tk.TclError:
            pass
        self.last_mtime, self.last_content = m, content
        self.update_footer()
        self.set_status("saved", "Atualizado pelo Drive")

    def tick(self):
        self._tick_n += 1
        try:
            self.check_external()
            if self._tick_n % RESCAN_EVERY_TICKS == 0:
                self.rescan()
        except Exception:
            pass
        self.root.after(TICK_MS, self.tick)

    def on_close(self):
        for job in (self.autosave_job, self._search_job, self._words_job):
            if job is not None:
                try:
                    self.root.after_cancel(job)
                except tk.TclError:
                    pass
        self.autosave_job = None
        if not self.leave_ok():
            return
        self.root.destroy()


def _enable_windows_dpi_awareness():
    """Evita que o Windows amplie a janela bitmap e deixe textos/contornos borrados."""
    if os.name != "nt":
        return
    try:
        import ctypes
        try:
            ctypes.windll.shcore.SetProcessDpiAwareness(2)
        except Exception:
            ctypes.windll.user32.SetProcessDPIAware()
    except Exception:
        pass


def main():
    _enable_windows_dpi_awareness()
    root = tk.Tk()
    NotebookApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
